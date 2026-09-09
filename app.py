"""Minimal, template-first PPTX renderer for the W&P Market Intelligence GPT.

The service deliberately edits a copy of Vorlage.pptx. It never builds a slide
from a blank canvas. Existing geometry, masters, footer, logo and chart slots
are treated as part of the public contract.
"""

from __future__ import annotations

import json
import os
import tempfile
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData


ROOT = Path(__file__).resolve().parent
TEMPLATE_PATH = ROOT / "Vorlage.pptx"
MANIFEST_PATH = ROOT / "template_manifest.json"
EXPECTED_WIDTH = 12_192_000
EXPECTED_HEIGHT = 6_858_000


def load_manifest() -> dict[str, Any]:
    return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))


MANIFEST = load_manifest()
TEXT_SLOTS = {str(k): int(v) for k, v in MANIFEST["text_slots"].items()}
CHART_SLOTS = {
    str(k): int(v["shape_id"]) for k, v in MANIFEST["chart_slots"].items()
}
CHART_TYPES = {
    str(k): str(v["template_type"]) for k, v in MANIFEST["chart_slots"].items()
}


class RenderError(ValueError):
    status = HTTPStatus.UNPROCESSABLE_ENTITY


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6:  # GROUP
            yield from iter_shapes(shape.shapes)


def find_shape(slide, shape_id: int):
    for shape in iter_shapes(slide.shapes):
        if int(shape.shape_id) == int(shape_id):
            return shape
    raise RenderError(f"Vorlagenfeld mit Shape-ID {shape_id} wurde nicht gefunden.")


def shape_geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def snapshot_geometry(slide) -> dict[int, tuple[int, int, int, int]]:
    return {int(s.shape_id): shape_geometry(s) for s in iter_shapes(slide.shapes)}


def set_shape_text(slide, shape_id: int, value: Any) -> None:
    shape = find_shape(slide, shape_id)
    if not getattr(shape, "has_text_frame", False):
        raise RenderError(f"Shape {shape_id} ist kein Textfeld.")
    text = str(value)
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text


def apply_text(slide, payload: dict[str, Any]) -> None:
    direct = {
        "title": TEXT_SLOTS["title"],
        "summary": TEXT_SLOTS["summary"],
        "notes_header": TEXT_SLOTS["notes_header"],
        "footnote": TEXT_SLOTS["footnote"],
    }
    for key, shape_id in direct.items():
        if key in payload and payload[key] is not None:
            set_shape_text(slide, shape_id, payload[key])

    section_titles = payload.get("section_titles", [])
    if section_titles:
        ids = [TEXT_SLOTS[f"chart_title_{i}"] for i in range(1, 5)]
        if len(section_titles) > len(ids):
            raise RenderError("Es sind höchstens vier Abschnittstitel vorgesehen.")
        for value, shape_id in zip(section_titles, ids):
            set_shape_text(slide, shape_id, value)

    for raw_id, value in payload.get("text_fields", {}).items():
        try:
            shape_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise RenderError(f"Ungültige Shape-ID: {raw_id}") from exc
        set_shape_text(slide, shape_id, value)


def remove_chart_slot(slide, slot: str) -> None:
    shape = find_shape(slide, CHART_SLOTS[slot])
    if not getattr(shape, "has_chart", False):
        raise RenderError(f"Slot {slot} enthält in der Vorlage kein Diagramm.")
    parent = shape._element.getparent()
    parent.remove(shape._element)


def replace_chart_data(slide, chart_spec: dict[str, Any]) -> None:
    slot = str(chart_spec.get("slot", ""))
    if slot not in CHART_SLOTS:
        raise RenderError(f"Unbekannter Diagramm-Slot: {slot}")
    if chart_spec.get("omit"):
        remove_chart_slot(slide, slot)
        return

    categories = chart_spec.get("categories")
    series = chart_spec.get("series")
    if not isinstance(categories, list) or not categories:
        raise RenderError(f"{slot}: categories muss eine nichtleere Liste sein.")
    if not isinstance(series, list) or not series:
        raise RenderError(f"{slot}: series muss eine nichtleere Liste sein.")
    if len(categories) > 24 or len(series) > 8:
        raise RenderError(f"{slot}: Diagrammdaten überschreiten die Vorlagengrenzen.")
    if CHART_TYPES[slot] == "doughnut" and len(series) != 1:
        raise RenderError(f"{slot}: Das Doughnut-Feld der Vorlage benötigt genau eine Serie.")

    chart_shape = find_shape(slide, CHART_SLOTS[slot])
    if not getattr(chart_shape, "has_chart", False):
        raise RenderError(f"{slot}: kein editierbares Diagramm in der Vorlage.")

    data = CategoryChartData()
    data.categories = [str(item) for item in categories]
    for item in series:
        if not isinstance(item, dict) or "name" not in item or "values" not in item:
            raise RenderError(f"{slot}: jede Serie benötigt name und values.")
        values = item["values"]
        if len(values) != len(categories):
            raise RenderError(f"{slot}: jede Serie muss gleich viele Werte wie Kategorien haben.")
        data.add_series(str(item["name"]), [float(v) for v in values])

    # replace_data keeps the chart type and the template formatting. The type
    # is intentionally fixed per slot to protect the one-to-one visual layout.
    chart_shape.chart.replace_data(data)


def apply_charts(slide, payload: dict[str, Any]) -> None:
    charts = payload.get("charts", [])
    if not isinstance(charts, list):
        raise RenderError("charts muss eine Liste sein.")
    for chart_spec in charts:
        if not isinstance(chart_spec, dict):
            raise RenderError("Jedes Diagramm muss ein Objekt sein.")
        replace_chart_data(slide, chart_spec)


def validate_layout(
    before: dict[int, tuple[int, int, int, int]],
    presentation: Presentation,
) -> None:
    if int(presentation.slide_width) != EXPECTED_WIDTH:
        raise RenderError("Die Folienbreite weicht von der Vorlage ab.")
    if int(presentation.slide_height) != EXPECTED_HEIGHT:
        raise RenderError("Die Folienhöhe weicht von der Vorlage ab.")
    after = snapshot_geometry(presentation.slides[0])
    for shape_id, geometry in before.items():
        if shape_id in after and after[shape_id] != geometry:
            raise RenderError(f"Layoutabweichung bei Shape {shape_id}.")


def render_presentation(payload: dict[str, Any]) -> bytes:
    if not TEMPLATE_PATH.exists():
        raise RenderError("Vorlage.pptx fehlt im Renderer-Repository.")
    if payload.get("template_id", "wandp-vorlage-v1") != MANIFEST["template_id"]:
        raise RenderError("Unbekannte template_id.")

    with tempfile.TemporaryDirectory(prefix="wandp-render-") as tmp:
        source = Path(tmp) / "source.pptx"
        output = Path(tmp) / "output.pptx"
        source.write_bytes(TEMPLATE_PATH.read_bytes())
        presentation = Presentation(str(source))
        if len(presentation.slides) != 1:
            raise RenderError("Die Vorlage muss genau eine Slide-Struktur enthalten.")
        if int(presentation.slide_width) != EXPECTED_WIDTH or int(presentation.slide_height) != EXPECTED_HEIGHT:
            raise RenderError("Die Vorlagengröße ist nicht kompatibel.")

        slide = presentation.slides[0]
        before = snapshot_geometry(slide)
        apply_text(slide, payload)
        apply_charts(slide, payload)
        presentation.save(str(output))

        # Reopen the result. This catches corrupt packages before the GPT sees one.
        checked = Presentation(str(output))
        validate_layout(before, checked)
        return output.read_bytes()


def openapi_document() -> dict[str, Any]:
    base_url = os.environ.get("RENDER_EXTERNAL_URL", "http://localhost:10000")
    return {
        "openapi": "3.1.0",
        "info": {
            "title": "W&P Template PPTX Renderer",
            "version": "1.0.0",
            "description": "Erzeugt eine PPTX ausschließlich auf Basis der festen W&P-Vorlage.",
        },
        "servers": [{"url": base_url}],
        "security": [{"bearerAuth": []}],
        "paths": {
            "/health": {"get": {"security": [], "responses": {"200": {"description": "OK"}}}},
            "/template-manifest": {"get": {"responses": {"200": {"description": "Vorlagenvertrag"}}}},
            "/render-pptx": {
                "post": {
                    "operationId": "renderPptx",
                    "summary": "Erzeugt eine PPTX auf Basis von Vorlage.pptx",
                    "requestBody": {
                        "required": True,
                        "content": {"application/json": {"schema": {"$ref": "#/components/schemas/RenderRequest"}}},
                    },
                    "responses": {
                        "200": {
                            "description": "Gerenderte PowerPoint-Datei",
                            "content": {"application/vnd.openxmlformats-officedocument.presentationml.presentation": {}},
                        },
                        "401": {"description": "Ungültiger API-Key"},
                        "413": {"description": "Anfrage zu groß"},
                        "422": {"description": "Vorlage oder Inhalt verletzt den Layoutvertrag"},
                    },
                }
            },
        },
        "components": {
            "securitySchemes": {"bearerAuth": {"type": "http", "scheme": "bearer"}},
            "schemas": {
                "Series": {
                    "type": "object",
                    "required": ["name", "values"],
                    "properties": {"name": {"type": "string"}, "values": {"type": "array", "items": {"type": "number"}}},
                },
                "Chart": {
                    "type": "object",
                    "required": ["slot"],
                    "properties": {
                        "slot": {"type": "string", "enum": ["chart_1", "chart_2", "chart_3", "chart_4"]},
                        "omit": {"type": "boolean"},
                        "categories": {"type": "array", "items": {"type": "string"}},
                        "series": {"type": "array", "items": {"$ref": "#/components/schemas/Series"}},
                    },
                },
                "RenderRequest": {
                    "type": "object",
                    "properties": {
                        "template_id": {"type": "string", "default": "wandp-vorlage-v1"},
                        "title": {"type": "string"},
                        "summary": {"type": "string"},
                        "notes_header": {"type": "string"},
                        "footnote": {"type": "string"},
                        "section_titles": {"type": "array", "items": {"type": "string"}, "maxItems": 4},
                        "text_fields": {"type": "object", "additionalProperties": {"type": "string"}},
                        "charts": {"type": "array", "items": {"$ref": "#/components/schemas/Chart"}},
                    },
                },
            },
        },
    }


def authorized(handler: BaseHTTPRequestHandler) -> bool:
    configured = os.environ.get("RENDERER_API_KEY")
    if not configured:
        # Local development remains usable. Render production should always set it.
        return True
    return handler.headers.get("Authorization", "") == f"Bearer {configured}"


class Handler(BaseHTTPRequestHandler):
    server_version = "WandPTemplateRenderer/1.0"

    def log_message(self, format: str, *args):  # noqa: A002
        print(format % args, flush=True)

    def send_json(self, status: int, value: Any) -> None:
        body = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):  # noqa: N802
        if self.path == "/health":
            return self.send_json(HTTPStatus.OK, {"status": "ok", "template": TEMPLATE_PATH.name})
        if self.path == "/openapi.json":
            return self.send_json(HTTPStatus.OK, openapi_document())
        if self.path == "/template-manifest":
            if not authorized(self):
                return self.send_json(HTTPStatus.UNAUTHORIZED, {"error": "invalid_proxy_key"})
            return self.send_json(HTTPStatus.OK, MANIFEST)
        return self.send_json(HTTPStatus.NOT_FOUND, {"error": "not_found"})

    def do_POST(self):  # noqa: N802
        if self.path != "/render-pptx":
            return self.send_json(HTTPStatus.NOT_FOUND, {"error": "not_found"})
        if not authorized(self):
            return self.send_json(HTTPStatus.UNAUTHORIZED, {"error": "invalid_renderer_key"})
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            length = 0
        if length <= 0 or length > 1_000_000:
            return self.send_json(HTTPStatus.REQUEST_ENTITY_TOO_LARGE, {"error": "request_too_large"})
        try:
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            if not isinstance(payload, dict):
                raise RenderError("Der Request muss ein JSON-Objekt sein.")
            body = render_presentation(payload)
        except RenderError as exc:
            return self.send_json(exc.status, {"error": "layout_contract_failed", "message": str(exc)})
        except Exception as exc:  # Keep internals out of the public API.
            print(f"render failure: {exc}", flush=True)
            return self.send_json(HTTPStatus.INTERNAL_SERVER_ERROR, {"error": "render_failed"})

        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.send_header("Content-Disposition", 'attachment; filename="market-analysis.pptx"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def main() -> None:
    port = int(os.environ.get("PORT", "10000"))
    server = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    print(f"W&P PPTX renderer listening on :{port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
